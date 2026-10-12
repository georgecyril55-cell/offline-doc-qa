import os
import pymupdf
import docx
import pandas as pd

folder_name = "Test_Documents"
os.makedirs(folder_name, exist_ok=True)

pdf_path = os.path.join(folder_name, "aerospace_handbook.pdf")
pdf_doc = pymupdf.open()
page = pdf_doc.new_page()
page.insert_text((50, 50), "Aerospace Handbook: The primary focus of this document is flight dynamics and parasite drag theory.")
pdf_doc.save(pdf_path)
pdf_doc.close()

docx_path = os.path.join(folder_name, "drag_calculations.docx")
word_doc = docx.Document()
word_doc.add_heading('Drag Calculation Report', 0)
word_doc.add_paragraph('The total calculated parasite drag coefficient (CD0) for the new test wing is exactly 0.045.')
word_doc.save(docx_path)

excel_path = os.path.join(folder_name, "project_budget.xlsx")
df = pd.DataFrame({
    "Project Phase": ["Design", "Prototyping", "Wind Tunnel Testing"],
    "Budget Allocated": ["$150,000", "$300,000", "$250,000"],
    "Status": ["Completed", "In Progress", "Pending"]
})
df.to_excel(excel_path, index=False)

print(f"Success! Dummy folder '{folder_name}' created with 3 test files.")